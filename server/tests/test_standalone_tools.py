import json
import struct
import tempfile
import unittest
from pathlib import Path

from docx import Document

from tools.altium_cfb_inspect import (
    ALTIUM_COORD_TO_MM,
    extract_board_bounds,
    iter_schematic_text,
)
from tools.generate_board_description_docx import generate


class FakeCFB:
    def __init__(self, streams):
        self.streams = streams

    def read_stream(self, path):
        return self.streams[path]


class AltiumCFBInspectTest(unittest.TestCase):
    def test_extract_board_bounds_converts_internal_coordinates_to_mm(self):
        header = b"V7_LAYER=1|"
        coordinates = (10.0, 20.0, 110.0, 20.0, 110.0, 70.0, 10.0, 70.0)
        stream = (
            struct.pack("<I", len(header))
            + header
            + struct.pack("<I", 4)
            + struct.pack("<8d", *coordinates)
        )
        bounds = extract_board_bounds(FakeCFB({"BoardRegions/Data": stream}))

        self.assertEqual(bounds.vertices, 4)
        self.assertAlmostEqual(bounds.width_mm, 100.0 * ALTIUM_COORD_TO_MM)
        self.assertAlmostEqual(bounds.height_mm, 50.0 * ALTIUM_COORD_TO_MM)

    def test_iter_schematic_text_prefers_utf8_text_field(self):
        record = "|Text=Fallback|%UTF8%Text=Строка~1две|".encode()
        stream = struct.pack("<I", len(record)) + record

        values = list(iter_schematic_text(FakeCFB({"FileHeader": stream})))

        self.assertEqual(values, ["Строка\nдве"])


class BoardDescriptionGeneratorTest(unittest.TestCase):
    def test_generate_uses_only_cli_supplied_template_data_and_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.docx"
            data = root / "description.json"
            output = root / "result.docx"

            template_doc = Document()
            template_doc.add_paragraph("OLD TITLE")
            template_doc.add_paragraph("old subtitle")
            template_doc.add_paragraph("01  Old section")
            table = template_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Parameter"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Old"
            table.cell(1, 1).text = "Old value"
            template_doc.sections[0].footer.paragraphs[0].text = "Old footer"
            template_doc.save(template)

            data.write_text(
                json.dumps(
                    {
                        "title": "ТЕХНИЧЕСКИЕ ХАРАКТЕРИСТИКИ",
                        "subtitle": "печатной платы TestBoard",
                        "footer_label": "TestBoard",
                        "sections": [
                            {
                                "number": "01",
                                "title": "Общие сведения",
                                "rows": [
                                    ["Назначение", "Испытательная плата"],
                                    ["Размер", "10 × 20 мм"],
                                ],
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            generate(template, data, output)

            result = Document(output)
            self.assertEqual(result.paragraphs[0].text, "ТЕХНИЧЕСКИЕ ХАРАКТЕРИСТИКИ")
            self.assertEqual(result.paragraphs[1].text, "печатной платы TestBoard")
            self.assertEqual(result.tables[0].cell(2, 1).text, "10 × 20 мм")
            footer_xml = result.sections[0].footer._element.xml
            self.assertIn("TestBoard", footer_xml)
            self.assertIn("PAGE", footer_xml)


if __name__ == "__main__":
    unittest.main()
