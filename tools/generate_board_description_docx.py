"""Generate a board-description DOCX from a Word template and JSON data.

The template supplies page geometry, styles, table formatting, and page breaks.
The JSON file supplies all board-specific text. No input or output paths are
embedded in the script.
"""

from __future__ import annotations

import argparse
import json
import sys
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.document import Document as DocumentObject
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # pragma: no cover - depends on caller environment
    raise SystemExit(
        "python-docx is required; install the project dependencies first"
    ) from exc


DEFAULT_COLORS = {
    "heading_number": "167FB5",
    "heading_text": "1D4059",
    "body_text": "334856",
    "alternate_row": "F2F6F8",
    "white": "FFFFFF",
    "footer": "607487",
}


@dataclass(frozen=True)
class SectionSpec:
    number: str
    title: str
    rows: list[tuple[str, str]]


@dataclass(frozen=True)
class DocumentSpec:
    title: str
    subtitle: str
    footer_label: str
    metadata_title: str
    metadata_subject: str
    sections: list[SectionSpec]
    colors: dict[str, str]
    left_column_dxa: int
    right_column_dxa: int


class SpecError(ValueError):
    """Raised when the JSON description does not match the expected schema."""


def require_string(value: Any, name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise SpecError(f"{name} must be a non-empty string")
    return value


def load_spec(path: Path) -> DocumentSpec:
    try:
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise SpecError(f"invalid JSON in {path}: {exc}") from exc
    if not isinstance(raw, dict):
        raise SpecError("the JSON root must be an object")

    title = require_string(raw.get("title"), "title")
    subtitle = require_string(raw.get("subtitle"), "subtitle")
    footer_label = require_string(raw.get("footer_label"), "footer_label")

    metadata = raw.get("metadata", {})
    if not isinstance(metadata, dict):
        raise SpecError("metadata must be an object")
    metadata_title = str(metadata.get("title", subtitle)).strip()
    metadata_subject = str(metadata.get("subject", "")).strip()

    colors = dict(DEFAULT_COLORS)
    custom_colors = raw.get("colors", {})
    if not isinstance(custom_colors, dict):
        raise SpecError("colors must be an object")
    for key, value in custom_colors.items():
        if key not in colors:
            raise SpecError(f"unsupported color key: {key}")
        text = str(value).lstrip("#").upper()
        if len(text) != 6 or any(char not in "0123456789ABCDEF" for char in text):
            raise SpecError(f"colors.{key} must be a six-digit RGB value")
        colors[key] = text

    layout = raw.get("layout", {})
    if not isinstance(layout, dict):
        raise SpecError("layout must be an object")
    left_column_dxa = int(layout.get("left_column_dxa", 3685))
    right_column_dxa = int(layout.get("right_column_dxa", 5783))
    if left_column_dxa <= 0 or right_column_dxa <= 0:
        raise SpecError("column widths must be positive integers")

    raw_sections = raw.get("sections")
    if not isinstance(raw_sections, list) or not raw_sections:
        raise SpecError("sections must be a non-empty array")
    sections: list[SectionSpec] = []
    for section_index, raw_section in enumerate(raw_sections, start=1):
        if not isinstance(raw_section, dict):
            raise SpecError(f"sections[{section_index}] must be an object")
        number = require_string(
            raw_section.get("number"), f"sections[{section_index}].number"
        )
        section_title = require_string(
            raw_section.get("title"), f"sections[{section_index}].title"
        )
        raw_rows = raw_section.get("rows")
        if not isinstance(raw_rows, list) or not raw_rows:
            raise SpecError(f"sections[{section_index}].rows must be a non-empty array")
        rows: list[tuple[str, str]] = []
        for row_index, row in enumerate(raw_rows, start=1):
            if (
                not isinstance(row, list)
                or len(row) != 2
                or not all(isinstance(item, str) for item in row)
            ):
                raise SpecError(
                    f"sections[{section_index}].rows[{row_index}] must be "
                    "a two-string array"
                )
            rows.append((row[0], row[1]))
        sections.append(SectionSpec(number, section_title, rows))

    return DocumentSpec(
        title=title,
        subtitle=subtitle,
        footer_label=footer_label,
        metadata_title=metadata_title,
        metadata_subject=metadata_subject,
        sections=sections,
        colors=colors,
        left_column_dxa=left_column_dxa,
        right_column_dxa=right_column_dxa,
    )


def set_font(run, *, size: Pt, bold: bool, color: str) -> None:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:cs"), "Aptos")
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    old = properties.find(qn("w:shd"))
    if old is not None:
        properties.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def replace_paragraph_text(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool,
    color: str,
    size: Pt,
    alignment=None,
) -> None:
    if not paragraph.runs:
        paragraph.add_run()
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)
    set_font(paragraph.runs[0], size=size, bold=bold, color=color)
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08


def set_cell_text(cell, text: str, *, bold: bool, color: str, header: bool) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(
        paragraph,
        text,
        bold=bold,
        color=color,
        size=Pt(9),
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", 85), ("bottom", 85), ("left", 125), ("right", 125)):
        side = margins.find(qn(f"w:{tag}"))
        if side is None:
            side = OxmlElement(f"w:{tag}")
            margins.append(side)
        side.set(qn("w:w"), str(value))
        side.set(qn("w:type"), "dxa")
    if header:
        set_cell_shading(cell, color)


def mark_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        properties.append(OxmlElement("w:tblHeader"))


def mark_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))
    fixed_height = properties.find(qn("w:trHeight"))
    if fixed_height is not None:
        properties.remove(fixed_height)


def resize_table(table, section: SectionSpec, spec: DocumentSpec) -> None:
    target_rows = 1 + len(section.rows)
    if len(table.rows) < 2:
        raise SpecError("every template table must contain a header and one body row")
    while len(table.rows) < target_rows:
        table._tbl.append(deepcopy(table.rows[1]._tr))
    while len(table.rows) > target_rows:
        table._tbl.remove(table.rows[-1]._tr)

    total_width = spec.left_column_dxa + spec.right_column_dxa
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(total_width))
    table_width.set(qn("w:type"), "dxa")
    grid_columns = table._tbl.tblGrid.gridCol_lst
    if len(grid_columns) != 2:
        raise SpecError("every template table must have exactly two columns")
    grid_columns[0].set(qn("w:w"), str(spec.left_column_dxa))
    grid_columns[1].set(qn("w:w"), str(spec.right_column_dxa))

    set_cell_text(
        table.cell(0, 0),
        "Параметр",
        bold=True,
        color=spec.colors["white"],
        header=True,
    )
    set_cell_shading(table.cell(0, 0), spec.colors["heading_text"])
    set_cell_text(
        table.cell(0, 1),
        "Значение",
        bold=True,
        color=spec.colors["white"],
        header=True,
    )
    set_cell_shading(table.cell(0, 1), spec.colors["heading_text"])
    set_cell_width(table.cell(0, 0), spec.left_column_dxa)
    set_cell_width(table.cell(0, 1), spec.right_column_dxa)
    mark_repeat_header(table.rows[0])

    for index, (label, value) in enumerate(section.rows, start=1):
        row = table.rows[index]
        mark_cant_split(row)
        fill = spec.colors["white"] if index % 2 else spec.colors["alternate_row"]
        for cell in row.cells:
            set_cell_shading(cell, fill)
        set_cell_width(row.cells[0], spec.left_column_dxa)
        set_cell_width(row.cells[1], spec.right_column_dxa)
        set_cell_text(
            row.cells[0],
            label,
            bold=True,
            color=spec.colors["heading_text"],
            header=False,
        )
        set_cell_text(
            row.cells[1],
            value,
            bold=False,
            color=spec.colors["body_text"],
            header=False,
        )


def heading_before_table(document: DocumentObject, table) -> Paragraph:
    element = table._tbl.getprevious()
    while element is not None:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document._body)
            if paragraph.text.strip():
                return paragraph
        element = element.getprevious()
    raise SpecError("could not find a heading paragraph before a template table")


def set_heading(paragraph: Paragraph, section: SectionSpec, spec: DocumentSpec) -> None:
    while len(paragraph.runs) < 2:
        paragraph.add_run()
    paragraph.runs[0].text = section.number
    paragraph.runs[1].text = "  " + section.title
    for run in paragraph.runs[2:]:
        run._element.getparent().remove(run._element)
    set_font(
        paragraph.runs[0], size=Pt(15), bold=True, color=spec.colors["heading_number"]
    )
    set_font(
        paragraph.runs[1], size=Pt(15), bold=True, color=spec.colors["heading_text"]
    )
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(5)


def nonempty_paragraphs_before_first_table(document: DocumentObject) -> list[Paragraph]:
    if not document.tables:
        raise SpecError("the template does not contain any tables")
    first_table = document.tables[0]._tbl
    result: list[Paragraph] = []
    for element in document._body._element:
        if element is first_table:
            break
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document._body)
            if paragraph.text.strip():
                result.append(paragraph)
    return result


def set_footer(document: DocumentObject, label: str, color: str) -> None:
    seen_parts: set[str] = set()
    for section in document.sections:
        footer = section.footer
        part_name = str(footer.part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        paragraph = (
            footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        )
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        prefix = paragraph.add_run(f"{label}  •  ")
        set_font(prefix, size=Pt(8), bold=False, color=color)

        begin_run = OxmlElement("w:r")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin_run.append(begin)
        paragraph._p.append(begin_run)

        instruction_run = OxmlElement("w:r")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        instruction_run.append(instruction)
        paragraph._p.append(instruction_run)

        separate_run = OxmlElement("w:r")
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run.append(separate)
        paragraph._p.append(separate_run)

        value = paragraph.add_run("1")
        set_font(value, size=Pt(8), bold=False, color=color)

        end_run = OxmlElement("w:r")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run.append(end)
        paragraph._p.append(end_run)


def generate(template: Path, data: Path, output: Path) -> None:
    spec = load_spec(data)
    document = Document(template)
    if len(document.tables) != len(spec.sections):
        raise SpecError(
            f"template contains {len(document.tables)} tables, but JSON defines "
            f"{len(spec.sections)} sections"
        )

    title_paragraphs = nonempty_paragraphs_before_first_table(document)
    if len(title_paragraphs) < 2:
        raise SpecError(
            "the template must have title and subtitle paragraphs before the first table"
        )
    replace_paragraph_text(
        title_paragraphs[0],
        spec.title,
        bold=True,
        color=spec.colors["heading_text"],
        size=Pt(22),
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    replace_paragraph_text(
        title_paragraphs[1],
        spec.subtitle,
        bold=False,
        color=spec.colors["heading_number"],
        size=Pt(14),
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    for table, section in zip(document.tables, spec.sections, strict=True):
        set_heading(heading_before_table(document, table), section, spec)
        resize_table(table, section, spec)

    set_footer(document, spec.footer_label, spec.colors["footer"])
    document.core_properties.title = spec.metadata_title
    document.core_properties.subject = spec.metadata_subject
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate a technical board-description DOCX from a template and JSON data."
    )
    parser.add_argument(
        "--template", required=True, type=Path, help="Source DOCX template"
    )
    parser.add_argument(
        "--data", required=True, type=Path, help="UTF-8 JSON description"
    )
    parser.add_argument(
        "--output", required=True, type=Path, help="Destination DOCX path"
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    args = build_parser().parse_args(argv)
    try:
        generate(args.template, args.data, args.output)
        print(args.output.resolve())
        return 0
    except (OSError, SpecError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
