"""Generate a board test-access DOCX from structured Altium inventory data."""

from __future__ import annotations

import argparse
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if __package__:
    from .board_test_localization import LANGUAGE_MODES, localized_text, ui_text
else:
    from board_test_localization import LANGUAGE_MODES, localized_text, ui_text

if __package__:
    from .generate_end_user_board_description_docx import (
        COLORS, SpecError, add_page_number, add_svg_picture, add_text,
        clear_paragraph_borders, fit_image, require_text, set_cant_split,
        set_cell_fill, set_cell_margins, set_cell_text, set_cell_width,
        set_keep_with_next, set_picture_alt, set_repeat_header, set_table_borders,
    )
else:
    from generate_end_user_board_description_docx import (
        COLORS, SpecError, add_page_number, add_svg_picture, add_text,
        clear_paragraph_borders, fit_image, require_text, set_cant_split,
        set_cell_fill, set_cell_margins, set_cell_text, set_cell_width,
        set_keep_with_next, set_picture_alt, set_repeat_header, set_table_borders,
    )


GENERATED_NET = re.compile(r"^Net[A-Za-z0-9_]+$")


def load_translation_catalog(path: Path | None) -> dict[str, Any]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict):
        raise SpecError("the translation catalog root must be an object")
    return payload


def require_localized(
    value: Any,
    field: str,
    language_mode: str,
    translations: dict[str, Any],
) -> str:
    text = localized_text(value, language_mode, translations=translations).strip()
    if not text:
        raise SpecError(f"{field} must be non-empty")
    return text


def require_verbatim_board_name(value: Any, board_name: str, field: str) -> None:
    if board_name not in str(value):
        raise SpecError(
            f"{field} must preserve board_name verbatim: {board_name!r}")


def resolve_path(spec_path: Path, value: Any, field: str) -> Path:
    path = Path(require_text(value, field))
    if not path.is_absolute():
        path = (spec_path.parent / path).resolve()
    if not path.exists():
        raise SpecError(f"missing file for {field}: {path}")
    return path


def parse_spec(
    path: Path,
    *,
    language_mode: str = "source",
    translations: dict[str, Any] | None = None,
) -> dict[str, Any]:
    raw = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(raw, dict):
        raise SpecError("the JSON root must be an object")
    translations = translations or {}
    board_name = require_text(raw.get("board_name"), "board_name")
    for field in ("title", "subtitle", "footer_label"):
        require_localized(raw.get(field), field, language_mode, translations)

    renders = raw.get("board_renders")
    maps = raw.get("test_maps")
    groups = raw.get("groups")
    if not isinstance(renders, list) or len(renders) != 2:
        raise SpecError("board_renders must contain top and bottom records")
    if not isinstance(maps, list) or len(maps) != 2:
        raise SpecError("test_maps must contain top and bottom records")
    if not isinstance(groups, list) or not groups:
        raise SpecError("groups must be a non-empty array")

    for index, render in enumerate(renders):
        require_localized(render.get("title"), f"board_renders[{index}].title", language_mode, translations)
        render["_image_path"] = resolve_path(path, render.get("image"), f"board_renders[{index}].image")
    for index, test_map in enumerate(maps):
        require_localized(test_map.get("title"), f"test_maps[{index}].title", language_mode, translations)
        test_map["_image_path"] = resolve_path(path, test_map.get("image"), f"test_maps[{index}].image")
        test_map["_fallback_path"] = resolve_path(path, test_map.get("fallback_image"), f"test_maps[{index}].fallback_image")
        if test_map["_image_path"].suffix.lower() != ".svg":
            raise SpecError(f"test_maps[{index}].image must be SVG")
        render_3d = test_map.get("render_3d")
        if render_3d is not None:
            if not isinstance(render_3d, dict):
                raise SpecError(f"test_maps[{index}].render_3d must be an object")
            require_localized(
                render_3d.get("title"), f"test_maps[{index}].render_3d.title",
                language_mode, translations)
            render_3d["_image_path"] = resolve_path(path, render_3d.get("image"), f"test_maps[{index}].render_3d.image")
            render_3d["_fallback_path"] = resolve_path(path, render_3d.get("fallback_image"), f"test_maps[{index}].render_3d.fallback_image")
            if render_3d["_image_path"].suffix.lower() != ".svg":
                raise SpecError(f"test_maps[{index}].render_3d.image must be SVG")

    seen: set[str] = set()
    for group_index, group in enumerate(groups):
        require_localized(group.get("title"), f"groups[{group_index}].title", language_mode, translations)
        items = group.get("items")
        if not isinstance(items, list) or not items:
            raise SpecError(f"groups[{group_index}].items must be non-empty")
        for item_index, item in enumerate(items):
            prefix = f"groups[{group_index}].items[{item_index}]"
            designator = require_text(item.get("designator"), f"{prefix}.designator")
            if designator in seen:
                raise SpecError(f"duplicate item designator: {designator}")
            seen.add(designator)
            require_localized(item.get("name"), f"{prefix}.name", language_mode, translations)
            require_text(item.get("side"), f"{prefix}.side")
            item["_image_path"] = resolve_path(path, item.get("image"), f"{prefix}.image")
            pins = item.get("pins")
            if not isinstance(pins, list) or not pins:
                raise SpecError(f"{prefix}.pins must be non-empty")
            for pin_index, pin in enumerate(pins):
                for field in ("pad", "pcb_net", "source_net"):
                    require_text(pin.get(field), f"{prefix}.pins[{pin_index}].{field}")
                if GENERATED_NET.match(str(pin["source_net"])):
                    raise SpecError(f"unresolved generated source net: {designator}.{pin['pad']}={pin['source_net']}")
    def localize(value: Any) -> str:
        return localized_text(value, language_mode, translations=translations)

    for field in ("title", "subtitle", "footer_label", "tagline", "summary_title", "summary_intro"):
        if field in raw:
            raw[field] = localize(raw[field])
    for row in raw.get("summary_rows", []):
        row[0], row[1] = localize(row[0]), localize(row[1])
    for render in renders:
        render["title"] = localize(render["title"])
        if "alt" in render:
            render["alt"] = localize(render["alt"])
    for test_map in maps:
        test_map["title"] = localize(test_map["title"])
        if "alt" in test_map:
            test_map["alt"] = localize(test_map["alt"])
        render_3d = test_map.get("render_3d")
        if render_3d:
            render_3d["title"] = localize(render_3d["title"])
            if "alt" in render_3d:
                render_3d["alt"] = localize(render_3d["alt"])
    for group in groups:
        group["title"] = localize(group["title"])
        if "intro" in group:
            group["intro"] = localize(group["intro"])
        for item in group["items"]:
            item["name"] = localize(item["name"])
            if item.get("description"):
                item["description"] = localize(item["description"])
    metadata = raw.get("metadata", {})
    for field in ("title", "subject"):
        if field in metadata:
            metadata[field] = localize(metadata[field])
    raw["_language_mode"] = language_mode
    for field in ("title", "subtitle"):
        require_verbatim_board_name(raw[field], board_name, field)
    return raw


def configure_document(document: Document, spec: dict[str, Any]) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor.from_string(COLORS["text"])
    normal.paragraph_format.space_after = Pt(3)
    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    clear_paragraph_borders(title)
    for level in (1, 2):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    add_text(footer, spec["footer_label"], size=7.8, bold=True, color=COLORS["muted"])
    footer_text = ui_text("footer", spec["_language_mode"], separator=" / ")
    add_text(footer, f"    |    {footer_text}    |    ", size=7.8, color=COLORS["muted"])
    add_page_number(footer)
    metadata = spec.get("metadata", {})
    document.core_properties.title = str(metadata.get("title", spec["subtitle"]))
    document.core_properties.subject = str(metadata.get("subject", "PCB test access and source nets"))


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph)
    add_text(paragraph, text, size=16 if level == 1 else 11, bold=True, color="000000")


def add_cover(document: Document, spec: dict[str, Any]) -> None:
    paragraph = document.add_paragraph(style="Title")
    clear_paragraph_borders(paragraph)
    paragraph.paragraph_format.space_before = Pt(45)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, spec["title"], size=27, bold=True, color="000000")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, spec["subtitle"], size=17, color=COLORS["blue"])
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        paragraph,
        str(spec.get("tagline", "Test points, configurable links, grounded contacts, and accessible signal holes")),
        size=10.5,
        color=COLORS["muted"],
    )
    document.add_page_break()


def add_summary(document: Document, spec: dict[str, Any]) -> None:
    add_heading(document, str(spec.get("summary_title", "Coverage and selection rules")), 1)
    paragraph = document.add_paragraph()
    add_text(paragraph, str(spec.get("summary_intro", "The inventory is derived from the open PCB and reports both the local PCB net and the traced source net.")), size=9)
    rows = spec.get("summary_rows", [])
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = (
        ui_text("scope", spec["_language_mode"], separator=" / "),
        ui_text("result_rule", spec["_language_mode"], separator=" / "),
    )
    for cell, width, label in zip(table.rows[0].cells, (3900, 9300), headers):
        set_cell_width(cell, width)
        set_cell_text(cell, label, bold=True, color=COLORS["white"], size=8.5)
        set_cell_fill(cell, COLORS["navy"])
    set_repeat_header(table.rows[0])
    for index, row_data in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        set_cell_width(row.cells[0], 3900)
        set_cell_width(row.cells[1], 9300)
        set_cell_text(row.cells[0], str(row_data[0]), bold=True, color=COLORS["navy"], size=8.4)
        set_cell_text(row.cells[1], str(row_data[1]), size=8.4)
        if index % 2:
            for cell in row.cells:
                set_cell_fill(cell, COLORS["pale"])
    document.add_page_break()


def add_full_page_image(document: Document, title: str, path: Path, alt: str) -> None:
    add_heading(document, title, 1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width, height = fit_image(path, max_width_cm=26.2, max_height_cm=15.0)
    shape = paragraph.add_run().add_picture(str(path), width=Cm(width), height=Cm(height))
    set_picture_alt(shape, alt)


def add_board_renders(document: Document, spec: dict[str, Any]) -> None:
    for render in spec["board_renders"]:
        add_full_page_image(document, str(render["title"]), render["_image_path"], str(render.get("alt", render["title"])))
        document.add_page_break()


def fit_svg_map(path: Path, *, max_width_cm: float, max_height_cm: float) -> tuple[float, float]:
    root = ET.parse(path).getroot()
    view_box = root.get("viewBox", "").replace(",", " ").split()
    if len(view_box) != 4:
        raise SpecError(f"map SVG has no four-value viewBox: {path}")
    width, height = float(view_box[2]), float(view_box[3])
    if width <= 0 or height <= 0:
        raise SpecError(f"map SVG has an invalid aspect ratio: {path}")
    scale = min(max_width_cm / width, max_height_cm / height)
    return width * scale, height * scale


def add_svg_map_page(
    document: Document,
    entry: dict[str, Any],
    *,
    language_mode: str,
    page_break_before: bool,
) -> None:
    add_heading(document, str(entry["title"]), 1)
    if page_break_before:
        document.paragraphs[-1].paragraph_format.page_break_before = True
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    max_height = 15.2 if language_mode == "en-zh" else 16.8
    width, height = fit_svg_map(
        entry["_image_path"], max_width_cm=26.7, max_height_cm=max_height)
    add_svg_picture(
        paragraph.add_run(),
        entry["_image_path"],
        entry["_fallback_path"],
        width=Cm(width),
        height=Cm(height),
        alt_text=str(entry.get("alt", entry["title"])),
    )


def add_test_maps(document: Document, spec: dict[str, Any]) -> None:
    for index, test_map in enumerate(spec["test_maps"]):
        add_svg_map_page(
            document, test_map, language_mode=spec["_language_mode"],
            page_break_before=bool(index))
        if test_map.get("render_3d"):
            add_svg_map_page(
                document, test_map["render_3d"], language_mode=spec["_language_mode"],
                page_break_before=True)


def add_pin_table(document: Document, item: dict[str, Any], language_mode: str) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = (1200, 3550, 4300, 1450)
    headers = tuple(
        ui_text(key, language_mode, separator=" / ")
        for key in ("pad", "pcb_net", "source_net", "hole_mm")
    )
    for cell, width, label in zip(table.rows[0].cells, widths, headers):
        set_cell_width(cell, width)
        set_cell_text(cell, label, bold=True, color=COLORS["white"], size=7.2)
        set_cell_fill(cell, COLORS["blue"])
    set_repeat_header(table.rows[0])
    for index, pin in enumerate(item["pins"]):
        row = table.add_row()
        set_cant_split(row)
        hole = float(pin.get("hole_mm") or 0.0)
        values = (pin["pad"], pin["pcb_net"], pin["source_net"], f"{hole:.2f}" if hole else "-")
        for cell, width, value in zip(row.cells, widths, values):
            set_cell_width(cell, width)
            set_cell_text(cell, str(value), size=7.2)
        if index % 2:
            for cell in row.cells:
                set_cell_fill(cell, COLORS["pale"])


def add_item(
    document: Document,
    item: dict[str, Any],
    *,
    language_mode: str,
    page_break_before: bool = False,
) -> None:
    add_heading(document, f"{item['designator']}  {item['name']}", 2)
    if page_break_before:
        document.paragraphs[-1].paragraph_format.page_break_before = True
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    set_keep_with_next(meta)
    side = ui_text(item["side"], language_mode, separator=" / ")
    add_text(meta, f"{side}  |  {item.get('footprint', '')}", size=7.6, bold=True, color=COLORS["muted"])
    description = str(item.get("description", "")).strip()
    if description:
        add_text(meta, f"  |  {description}", size=7.6, color=COLORS["muted"])

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    set_keep_with_next(paragraph)
    width, height = fit_image(item["_image_path"], max_width_cm=9.2, max_height_cm=5.2)
    shape = paragraph.add_run().add_picture(str(item["_image_path"]), width=Cm(width), height=Cm(height))
    set_picture_alt(shape, ui_text("crop_alt", language_mode, designator=item["designator"]))
    add_pin_table(document, item, language_mode)


def add_groups(document: Document, spec: dict[str, Any]) -> None:
    for group_index, group in enumerate(spec["groups"]):
        add_heading(document, str(group["title"]), 1)
        document.paragraphs[-1].paragraph_format.page_break_before = True
        intro = str(group.get("intro", "")).strip()
        if intro:
            paragraph = document.add_paragraph()
            add_text(paragraph, intro, size=8.5, color=COLORS["muted"])
        for index, item in enumerate(group["items"]):
            start_page = bool(
                index
                and (
                    index % 2 == 0
                    or "through-hole-signal" in item.get("categories", [])
                    or item.get("designator") == "SB4"
                )
            )
            add_item(
                document, item, language_mode=spec["_language_mode"],
                page_break_before=start_page)


def build_document(spec: dict[str, Any], output: Path) -> None:
    document = Document()
    configure_document(document, spec)
    add_cover(document, spec)
    add_summary(document, spec)
    add_board_renders(document, spec)
    add_test_maps(document, spec)
    add_groups(document, spec)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--language-mode", choices=LANGUAGE_MODES, default="source")
    parser.add_argument("--translations", type=Path)
    args = parser.parse_args()
    if args.data.resolve() == args.output.resolve():
        raise SystemExit("input and output paths must be different")
    try:
        translations = load_translation_catalog(
            args.translations.resolve() if args.translations else None)
        spec = parse_spec(
            args.data.resolve(), language_mode=args.language_mode,
            translations=translations)
        build_document(spec, args.output.resolve())
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
